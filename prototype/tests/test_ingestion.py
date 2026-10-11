from datetime import date
from io import BytesIO
from pathlib import Path
import json
import pytest
from docx import Document
from openpyxl import Workbook
from ingestion import extract_notes,prepare_sources,prepare_metrics,excel_sheets
from schemas import SourcePack
from workflow import build_report,load_sample,input_fingerprint

SAMPLES=Path(__file__).resolve().parents[1]/'sample_data/meeting_files'


def test_transcript_retains_speaker_names_and_timestamps():
    for name in ['partner-call.vtt','vendor-call.srt']:
        raw=(SAMPLES/name).read_bytes()
        text=extract_notes(name,raw)
        assert text==raw.decode().strip()
        assert '00:00:' in text
    assert 'Luca (Swisscom)' in extract_notes('partner-call.vtt',(SAMPLES/'partner-call.vtt').read_bytes())


def test_word_paragraphs_and_tables_keep_their_order():
    doc=Document();doc.add_paragraph('Before the table')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Action';table.cell(0,1).text='Still open'
    doc.add_paragraph('After the table');buffer=BytesIO();doc.save(buffer)
    text=extract_notes('notes.docx',buffer.getvalue())
    assert text.index('Before')<text.index('Action | Still open')<text.index('After')


def test_excel_matches_csv_calculations_and_preserves_row_provenance():
    raw=(SAMPLES/'platform-metrics.xlsx').read_bytes()
    assert excel_sheets(raw)==['Metrics','Read me']
    normalized,origin=prepare_metrics('platform-metrics.xlsx',raw,'Metrics')
    csv,notes=load_sample('normal')
    assert build_report(normalized,notes)['metrics']==build_report(csv,notes)['metrics']
    assert len(origin['conversions'])==4
    assert origin['row_map'][0]=={'normalized_csv_row':2,'original_row':2}
    assert origin['worksheet']=='Metrics'


def test_excel_formulas_are_not_evaluated_or_silently_treated_as_missing():
    wb=Workbook();ws=wb.active;ws.title='Metrics';ws.append(['Week start','Cohort','Metric','Value','Unit']);ws.append([date(2026,9,28),'MEP','availability','=99+0.9','%']);b=BytesIO();wb.save(b)
    with pytest.raises(ValueError,match='formula'): prepare_metrics('metrics.xlsx',b.getvalue(),'Metrics')


def test_metrics_extra_columns_and_blank_rows_have_explicit_provenance():
    raw=b'Week start,Cohort,Metric,Value,Unit,Comment\n2026-09-28,MEP,availability,99.9%,%,Verified\n,,,,,\n2026-09-21,MEP,availability,99.8,%,Previous\n'
    normalized,origin=prepare_metrics('metrics.csv',raw)
    assert origin['ignored_columns']==['Comment']
    assert origin['row_map'][1]=={'normalized_csv_row':3,'original_row':4}
    assert b'99.9,%' in normalized


def test_pasted_notes_need_no_json_and_do_not_get_an_invented_meeting_date():
    sources,manifest=prepare_sources([], 'BD still needs to confirm the acceptable analytics delay.',date(2026,9,28))
    pack=SourcePack.model_validate_json(sources)
    assert pack.sources[0].date is None
    assert pack.sources[0].title.startswith('Pasted notes.txt')
    assert manifest[0]['chunks'][0]['source_id']==pack.sources[0].id
    r=build_report(load_sample('normal')[0],sources,intake={'notes':manifest})
    assert r['narrative']['findings'][0]['category']=='source note'
    assert any(i['code']=='undated_sources' for i in r['issues'])


def test_long_notes_are_split_without_changing_or_dropping_nonwhitespace():
    text=('Maya: The action is still open.\nKen: We have not agreed a fix date.\n'*160).strip()
    sources,manifest=prepare_sources([('transcript.txt',text.encode())],'',date(2026,9,28))
    pack=SourcePack.model_validate_json(sources)
    assert len(pack.sources)>1
    assert ''.join(''.join(s.text.split()) for s in pack.sources)==''.join(text.split())
    for s,m in zip(pack.sources,manifest[0]['chunks']):
        assert text[m['extracted_text_start']:m['extracted_text_end']]==s.text
        assert len(s.text)<=6000


def test_oversized_notes_are_rejected_not_summarized_or_truncated():
    with pytest.raises(ValueError,match='40,000'): prepare_sources([], 'x'*40001,date(2026,9,28))


def test_source_file_changes_invalidate_input_identity_even_with_same_extracted_text():
    a,_=prepare_sources([('notes.txt',b'Notes for this week are unchanged.')],'',date(2026,9,28))
    assert input_fingerprint(b'a',a,{'sha':'file-one'})!=input_fingerprint(b'a',a,{'sha':'file-two'})


def test_empty_image_only_word_document_is_rejected():
    doc=Document();b=BytesIO();doc.save(b)
    with pytest.raises(ValueError,match='readable text'):extract_notes('scan.docx',b.getvalue())


def test_wrong_worksheet_and_non_monday_fail_clearly():
    with pytest.raises(ValueError,match='worksheet'):prepare_metrics('metrics.xlsx',(SAMPLES/'platform-metrics.xlsx').read_bytes(),'Missing')
    with pytest.raises(ValueError,match='Monday'):prepare_sources([], 'Some ordinary meeting notes.',date(2026,9,29))


def pdf_bytes(pages=('Vendor update: maintenance is proposed, not approved.',),password=None,padding=0):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter()
    for text in pages:
        page=writer.add_blank_page(width=612,height=792)
        if text:
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
            stream=DecodedStreamObject()
            stream.set_data(('BT /F1 12 Tf 50 740 Td ('+text+') Tj ET').encode())
            page[NameObject('/Contents')]=stream
    if padding: writer.add_metadata({'/Subject':'x'*padding})
    if password: writer.encrypt(password)
    result=BytesIO();writer.write(result);return result.getvalue()


def test_pdf_text_and_page_markers_reach_report_with_source_provenance():
    raw=pdf_bytes(('Maintenance is proposed, not approved.','No outage or delivery date is confirmed.'))
    pack,manifest=prepare_sources([('vendor.pdf',raw)],'',date(2026,9,28))
    report=build_report(load_sample('normal')[0],pack,intake={'notes':manifest})
    text=report['pack']['sources'][0]['text']
    assert '[PDF page 1]' in text and '[PDF page 2]' in text
    assert 'No outage or delivery date is confirmed.' in text
    assert manifest[0]['filename']=='vendor.pdf'
    assert not any(i['severity']=='error' for i in report['issues'])


@pytest.mark.parametrize('raw,message',[
    (pdf_bytes(('Readable first page.',None)),'page 2 has no readable text'),
    (pdf_bytes(password='test-password'),'Password-protected'),
    (b'not a PDF','Cannot read'),
    (pdf_bytes(tuple('A readable page.' for _ in range(51))),'1–50 pages'),
])
def test_pdf_rejects_unsupported_content_without_silent_partial_ingestion(raw,message):
    with pytest.raises(ValueError,match=message): extract_notes('document.pdf',raw)


def test_larger_files_are_accepted_but_ten_mb_limit_is_enforced():
    from ingestion import MAX_FILE_BYTES
    raw=pdf_bytes(padding=2_100_000)
    assert len(raw)>2_000_000
    assert 'maintenance is proposed' in extract_notes('large.pdf',raw)
    with pytest.raises(ValueError,match='10 MB'):
        extract_notes('oversized.pdf',b'x'*(MAX_FILE_BYTES+1))


def test_excel_auto_detection_uses_columns_not_sheet_name_or_position():
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO((SAMPLES/'platform-metrics.xlsx').read_bytes()))
    wb['Metrics'].title='Platform export September'
    wb.active=1
    wb.move_sheet('Read me',offset=-1)
    raw=BytesIO();wb.save(raw);wb.close()
    normalized,origin=prepare_metrics('export.xlsx',raw.getvalue())
    assert origin['worksheet']=='Platform export September'
    csv,notes=load_sample('normal')
    assert build_report(normalized,notes)['metrics']==build_report(csv,notes)['metrics']


def test_excel_auto_detection_blocks_ambiguous_sheets():
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO((SAMPLES/'platform-metrics.xlsx').read_bytes()))
    wb.copy_worksheet(wb['Metrics']).title='Other cohort'
    raw=BytesIO();wb.save(raw);wb.close()
    with pytest.raises(ValueError,match='More than one sheet'):
        prepare_metrics('ambiguous.xlsx',raw.getvalue())


def test_excel_auto_detection_requires_the_known_column_contract():
    wb=Workbook();wb.active.title='Metrics'
    wb.active.append(['Customers','Sales']);wb.active.append([100,200])
    raw=BytesIO();wb.save(raw);wb.close()
    with pytest.raises(ValueError,match='No platform data sheet found'):
        prepare_metrics('unknown.xlsx',raw.getvalue())
