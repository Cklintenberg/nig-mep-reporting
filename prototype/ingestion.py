"""Turn normal office files into validated internal inputs; no model call or silent truncation."""
import csv
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader
from schemas import Source, SourcePack

MAX_FILE_BYTES=10 * 1024 * 1024
REQUIRED=('week_start','cohort','metric','value','unit')
ALIASES={'week':'week_start','week_start':'week_start','week_starting':'week_start','reporting_week':'week_start',
         'cohort':'cohort','metric':'metric','value':'value','unit':'unit'}


def _check_size(raw):
    if len(raw)>MAX_FILE_BYTES: raise ValueError('Each upload must be no larger than 10 MB.')


def _pdf_text(filename, raw):
    """Read text pages only; do not silently discard pages needing OCR."""
    _check_size(raw)
    try:
        reader=PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
        if not 0 < len(reader.pages) <= 50:
            raise ValueError('Use a PDF with 1–50 pages for this report.')
        parts=[];stream_bytes=0
        for number,page in enumerate(reader.pages,start=1):
            contents=page.get_contents()
            stream_bytes+=len(contents.get_data()) if contents else 0
            if stream_bytes>20_000_000:
                raise ValueError('This PDF is too complex to read here. Export a smaller document or paste the relevant text.')
            text=(page.extract_text() or '').strip()
            if not text:
                raise ValueError(f'{filename}: page {number} has no readable text. Scanned pages need OCR first; remove blank pages before uploading.')
            parts.append(f'[PDF page {number}]\n{text}')
            if sum(len(part) for part in parts)>40_000:
                raise ValueError('The PDF exceeds 40,000 characters. Upload only the relevant pages; nothing has been truncated.')
        return '\n\n'.join(parts)
    except ValueError:
        raise
    except Exception:
        raise ValueError(f'Cannot read {filename}. Export it again as a text-based PDF or paste its text.') from None


def _check_office_archive(raw):
    _check_size(raw)
    try:
        with ZipFile(BytesIO(raw)) as archive:
            files=archive.infolist()
            if len(files)>2000 or sum(f.file_size for f in files)>20_000_000:
                raise ValueError('Office file expands beyond the prototype limit. Use a smaller export.')
            if any(f.flag_bits & 1 for f in files): raise ValueError('Password-protected files are not supported.')
    except BadZipFile:
        raise ValueError('This is not a readable DOCX/XLSX file. Export it again without a password.') from None


def _decode(raw):
    _check_size(raw)
    try:
        return raw.decode('utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
    except UnicodeError:
        raise ValueError('Text files must use UTF-8 or UTF-16 with a byte-order mark. You can also paste the notes.') from None


def extract_notes(filename, raw):
    suffix=Path(filename).suffix.lower()
    if suffix in {'.txt','.vtt','.srt'}:
        text=_decode(raw)
    elif suffix=='.pdf':
        text=_pdf_text(filename,raw)
    elif suffix=='.docx':
        _check_office_archive(raw)
        try:
            doc=Document(BytesIO(raw))
            blocks=[]
            for item in doc.iter_inner_content():
                if isinstance(item,Table):
                    blocks.extend(' | '.join(cell.text for cell in row.cells) for row in item.rows)
                else: blocks.append(item.text)
            text='\n\n'.join(blocks)
        except Exception:
            raise ValueError(f'Cannot extract text from {filename}. Try a DOCX export or paste the notes.') from None
    else:
        raise ValueError('Notes must be TXT, VTT, SRT, DOCX or a text-based PDF, or pasted as text. Images, audio and legacy DOC files are not supported.')
    text=text.replace('\r\n','\n').replace('\r','\n').strip()
    if len(text)<8: raise ValueError(f'{filename} contains too little readable text. Scans and images need transcription or OCR first.')
    if '\x00' in text: raise ValueError(f'{filename} is not readable plain text.')
    return text


def _parts(text, limit=5800):
    """Bounded contiguous excerpts with offsets in the extracted text, not summaries."""
    start=0
    while start<len(text):
        end=min(start+limit,len(text))
        if end<len(text):
            # Prefer complete paragraphs, lines or words; never silently drop content.
            cuts=[text.rfind(sep,start+limit//2,end) for sep in ['\n\n','\n',' ']]
            cut=next((c for c in cuts if c>=0),-1)
            if cut>=0: end=cut
            if len(text)-end<8: end=len(text)-8
        fragment=text[start:end]
        clean=fragment.strip()
        if clean:
            offset=start+len(fragment)-len(fragment.lstrip())
            yield clean,offset,offset+len(clean)
        start=end


def prepare_sources(files, pasted_notes, week_start):
    if len(files)>12: raise ValueError('Upload at most 12 note files per weekly report.')
    records=[]; manifests=[]
    inputs=[(Path(name).name,raw,extract_notes(name,raw)) for name,raw in files]
    if pasted_notes.strip():
        raw=pasted_notes.encode('utf-8')
        inputs.append(('Pasted notes.txt',raw,extract_notes('Pasted notes.txt',raw)))
    if not inputs: raise ValueError('Upload a transcript or document, or paste your meeting notes.')
    if sum(len(t) for _,_,t in inputs)>40_000:
        raise ValueError('The combined notes exceed 40,000 characters. Choose the material relevant to this week; nothing has been truncated.')
    for index,(name,raw,text) in enumerate(inputs,start=1):
        chunks=[]
        for part,(fragment,start,end) in enumerate(_parts(text),start=1):
            if len(fragment)<8:
                raise ValueError(f'{name} produced a fragment that is too short. Add a little context or combine it with another note.')
            source_id=f'DOC{index:02d}-{part:02d}'
            records.append(Source(id=source_id,kind='context',date=None,title=f'{name[:125]} · part {part}',text=fragment))
            chunks.append({'source_id':source_id,'extracted_text_start':start,'extracted_text_end':end})
        manifests.append({'filename':name,'sha256':sha256(raw).hexdigest(),'extracted_characters':len(text),'chunks':chunks,
                          'date_provenance':'Not inferred; calendar meeting date was not provided.'})
    pack=SourcePack(label='Weekly report from uploaded meeting material',week_start=week_start,sources=records)
    return pack.model_dump_json().encode(),manifests


def excel_sheets(raw):
    _check_office_archive(raw)
    try:
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=False)
        names=wb.sheetnames
        wb.close()
        return names
    except Exception:
        raise ValueError('Cannot read this Excel workbook. Use an unencrypted .xlsx export.') from None


def _header(value):
    return re.sub(r'[^a-z0-9]+','_',str(value or '').strip().lower()).strip('_')


def prepare_metrics(filename, raw, sheet=None):
    """Normalize a defined metric table; never guess business meanings or evaluate formulas."""
    suffix=Path(filename).suffix.lower();_check_size(raw)
    if suffix=='.xlsx':
        _check_office_archive(raw)
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=False)
        try:
            if sheet is None:
                if len(wb.worksheets)>50:
                    raise ValueError('Use a workbook with at most 50 sheets.')
                candidates=[]
                for candidate in wb.worksheets:
                    first=next(candidate.iter_rows(min_row=1,max_row=1,max_col=50,values_only=True),())
                    headers=[ALIASES.get(_header(v),_header(v)) for v in first]
                    if all(headers.count(key)==1 for key in REQUIRED):
                        candidates.append(candidate.title)
                if not candidates:
                    raise ValueError('No platform data sheet found. Use the first-row headings: Week start, Cohort, Metric, Value, Unit.')
                if len(candidates)>1:
                    raise ValueError('More than one sheet matches the platform data format. Upload a workbook containing just the intended data sheet, or export that sheet as CSV.')
                sheet=candidates[0]
            if sheet not in wb.sheetnames: raise ValueError('Choose the Excel worksheet containing the metrics.')
            ws=wb[sheet]
            if (ws.max_row or 0)>501 or (ws.max_column or 0)>50: raise ValueError('Use a table with at most 500 data rows and 50 columns; remove unused formatted rows/columns.')
            from itertools import islice
            rows=list(islice(ws.iter_rows(max_col=min(ws.max_column or 50,50)),502))
            if len(rows)>501: raise ValueError("Use at most 500 data rows.")
            values=[[c.value for c in row] for row in rows]
            formats=[[c.number_format for c in row] for row in rows]
            formulas=[[c.data_type=='f' for c in row] for row in rows]
        finally: wb.close()
    elif suffix=='.csv':
        try: values=list(csv.reader(StringIO(_decode(raw))))
        except csv.Error: raise ValueError('Cannot read this CSV. Export a comma-separated table.') from None
        if len(values)>501: raise ValueError('Use at most 500 data rows.')
        formats=[['']*len(r) for r in values];formulas=[[False]*len(r) for r in values]
    else: raise ValueError('Platform figures must be CSV or XLSX.')
    if not values: raise ValueError('The metrics table is empty.')
    headers=[ALIASES.get(_header(v),_header(v)) for v in values[0]]
    if any(headers.count(key)!=1 for key in REQUIRED):
        raise ValueError('The first row needs these column headings: Week start, Cohort, Metric, Value, Unit. The download examples show the format; arbitrary dashboards and pivot tables are not supported.')
    positions={key:headers.index(key) for key in REQUIRED}
    output=StringIO();writer=csv.writer(output,lineterminator='\n');writer.writerow(REQUIRED)
    row_map=[];conversions=[]
    for idx,row in enumerate(values[1:],start=1):
        if all(v is None or str(v).strip()=='' for v in row): continue
        if len(row)!=len(headers): raise ValueError(f'Row {idx+1} has a different number of cells than the header.')
        normalized={}
        for key,pos in positions.items():
            value=row[pos]
            if formulas[idx][pos]: raise ValueError(f'Row {idx+1} contains a formula. Export calculated values before uploading; formulas are not executed.')
            if key=='week_start' and isinstance(value,(date,datetime)): value=value.strftime('%Y-%m-%d')
            normalized[key]='' if value is None else str(value).strip()
        valuepos=positions['value']
        if normalized['unit']=='%' and isinstance(row[valuepos],(float,int)) and '%' in formats[idx][valuepos]:
            normalized['value']=str(row[valuepos]*100)
            conversions.append(f'Row {idx+1}: Excel percentage converted from its stored fraction to the 0–100 percentage scale.')
        elif normalized['unit']=='%' and normalized['value'].endswith('%'):
            normalized['value']=normalized['value'][:-1].strip()
        writer.writerow([normalized[key] for key in REQUIRED])
        row_map.append({'normalized_csv_row':len(row_map)+2,'original_row':idx+1})
    if not row_map: raise ValueError('The selected metrics table contains no data rows.')
    provenance={'filename':Path(filename).name,'sha256':sha256(raw).hexdigest(),'worksheet':sheet if suffix=='.xlsx' else None,
                'row_map':row_map,'conversions':conversions,'ignored_columns':[str(v) for v,h in zip(values[0],headers) if h not in REQUIRED]}
    return output.getvalue().encode(),provenance
